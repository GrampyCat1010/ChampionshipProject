import base64
import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from mediapodborka.content import extract, normalize_url, source_matches, public_addresses
from mediapodborka.db import connect
from mediapodborka.jobs import run_search
from mediapodborka.yandex import parse_analysis, parse_search_response, ServiceError

TEXT = 'В Красноярске завершено благоустройство набережной Енисея. Городская администрация сообщила об открытии нового пешеходного участка, освещении и доступной среде для жителей.'


def start(client, headers, payload):
    response = client.post('/api/searches', json=payload, headers=headers)
    assert response.status_code == 202, response.json
    return response.json['id']


def configure_pipeline(monkeypatch, candidates, publication=None, analysis=None):
    monkeypatch.setattr('mediapodborka.jobs.web_search', lambda *a, **kw: candidates)
    monkeypatch.setattr('mediapodborka.jobs.fetch_publication', lambda url: {'title': 'Набережная Енисея', 'original_text': TEXT, 'published_at': '2026-10-12', 'source': 'Редакция', 'content_hash': 'same-content', 'resolved_url': url, **(publication or {})})
    monkeypatch.setattr('mediapodborka.jobs.analyze', lambda *a: {'relevance': 'relevant', 'reason': 'Совпадают место и событие', 'evidence': TEXT[:60], 'region_reason': 'Красноярск прямо назван', 'verification': 'insufficient', **(analysis or {})})


def test_home_csrf_and_no_key_leaks(client):
    page = client.get('/')
    assert page.status_code == 200
    assert 'Content-Security-Policy' in page.headers
    assert b'YANDEX_CLOUD_API_KEY' not in page.data
    assert client.post('/api/searches', json={}).status_code == 403
    assert client.get('/.env').status_code == 404
    assert 'api_key' not in json.dumps(client.get('/api/status').json)


def test_source_crud_duplicate_and_session_isolation(app, client, headers):
    source = {'name': 'Редакция', 'url': 'https://vk.com/club123?utm_source=test', 'region': 'Красноярский край'}
    created = client.post('/api/sources', json=source, headers=headers)
    assert created.status_code == 201
    source_id = created.json['id']
    assert client.post('/api/sources', json=source, headers=headers).status_code == 409
    other = app.test_client()
    assert other.get('/api/sources').json['sources'] == []
    assert client.put('/api/sources/' + source_id, json={**source, 'name': 'Изменено'}, headers=headers).status_code == 200
    assert client.get('/api/sources').json['sources'][0]['name'] == 'Изменено'
    assert client.delete('/api/sources/' + source_id, headers=headers).status_code == 200
    assert client.get('/api/sources').json['sources'] == []


@pytest.mark.parametrize('change', [ {'date_from': '2026-11-01'}, {'date_from': 'bad'}, {'platforms': []}, {'platforms': ['invalid']}, {'query': 'a'}, {'mode': 'sources'}, {'query_type': 'bad'}, {'date_from': '2024-01-01'}, {'source_ids': 'bad'} ])
def test_validation(client, headers, search_payload, change):
    assert client.post('/api/searches', json={**search_payload, **change}, headers=headers).status_code == 400


def test_pipeline_dedup_and_export_selection(app, client, headers, search_payload, monkeypatch):
    candidates = [
        {'url': 'https://news.example.com/article/1?utm_source=a', 'title': 'Один', 'snippet': 'Выдержка'},
        {'url': 'https://news.example.com/article/1?utm_source=b', 'title': 'Повтор', 'snippet': 'Выдержка'},
        {'url': 'https://second.example.com/article/2', 'title': 'Перепечатка', 'snippet': 'Выдержка'},
    ]
    configure_pipeline(monkeypatch, candidates)
    sid = start(client, headers, search_payload)
    run_search(app.config, sid)
    result = client.get('/api/searches/' + sid).json
    assert result['search']['state'] == 'completed'
    assert len(result['publications']) == 2
    assert all(p['placements'] == 2 for p in result['publications'])
    p = result['publications'][0]
    assert client.patch(f'/api/searches/{sid}/publications/{p["id"]}', json={'included': False}, headers=headers).status_code == 200
    export = client.get(f'/api/searches/{sid}/export')
    assert export.status_code == 200
    wb = load_workbook(BytesIO(export.data))
    assert wb['Публикации'].max_row == 2
    assert wb['Публикации'].cell(2, 7).hyperlink.target == result['publications'][1]['url']
    assert wb['Условия поиска'].cell(2, 2).value == search_payload['query']
    with connect(app.config['DATABASE']) as db:
        db.execute('UPDATE publications SET title=? WHERE included=1', ('=HYPERLINK("evil")',))
    wb = load_workbook(BytesIO(client.get(f'/api/searches/{sid}/export').data))
    assert wb['Публикации']['D2'].data_type != 'f'
    other = app.test_client()
    assert other.get('/api/searches/' + sid).status_code == 404
    assert other.get(f'/api/searches/{sid}/export').status_code == 404


@pytest.mark.parametrize('published, expected', [(None, 'relevant'), ('2026-09-01', 'irrelevant')])
def test_unknown_and_outside_dates_not_included(app, client, headers, search_payload, monkeypatch, published, expected):
    configure_pipeline(monkeypatch, [{'url': 'https://news.example.com/a', 'title': 'Статья', 'snippet': ''}], {'published_at': published})
    sid = start(client, headers, search_payload)
    run_search(app.config, sid)
    pubs = client.get('/api/searches/' + sid).json['publications']
    if published:
        assert pubs == []  # Out-of-period publications are now excluded completely.
    else:
        assert pubs[0]['relevance'] == expected
        assert pubs[0]['included'] == 0
    assert client.get(f'/api/searches/{sid}/export').status_code == 400


def test_unavailable_original_and_model_failure(app, client, headers, search_payload, monkeypatch):
    configure_pipeline(monkeypatch, [{'url': 'https://news.example.com/a', 'title': 'Статья', 'snippet': 'Это только сниппет'}])
    def fail(*args):
        raise ValueError('Оригинал недоступен')
    monkeypatch.setattr('mediapodborka.jobs.fetch_publication', fail)
    sid = start(client, headers, search_payload)
    run_search(app.config, sid)
    result = client.get('/api/searches/' + sid).json
    assert result['search']['state'] == 'partial'
    assert result['publications'] == []  # A snippet cannot establish an exact match.


def test_empty_vs_failed_search(app, client, headers, search_payload, monkeypatch):
    configure_pipeline(monkeypatch, [])
    sid = start(client, headers, search_payload)
    run_search(app.config, sid)
    assert client.get('/api/searches/' + sid).json['search']['state'] == 'completed'
    def fail(*args):
        raise ServiceError('Нет доступа')
    monkeypatch.setattr('mediapodborka.jobs.web_search', fail)
    sid2 = start(client, headers, search_payload)
    run_search(app.config, sid2)
    assert client.get('/api/searches/' + sid2).json['search']['state'] == 'failed'


def test_source_mode_boundaries(app, client, headers, search_payload, monkeypatch):
    source = client.post('/api/sources', json={'name': 'Наш источник', 'url': 'https://vk.com/club123', 'region': 'Красноярск'}, headers=headers).json
    configure_pipeline(monkeypatch, [{'url': 'https://vk.com/wall-123_1', 'title': 'Наш', 'snippet': ''}, {'url': 'https://vk.com/wall-999_1', 'title': 'Чужой', 'snippet': ''}])
    sid = start(client, headers, {**search_payload, 'mode': 'sources', 'source_ids': [source['id']]})
    run_search(app.config, sid)
    pubs = client.get('/api/searches/' + sid).json['publications']
    assert len(pubs) == 1 and pubs[0]['url'] == 'https://vk.com/wall-123_1'


def test_analysis_evidence_must_exist():
    response = {'relevance': 'relevant', 'reason': 'Да', 'evidence': 'Вымышленная цитата, которой не было в оригинальном тексте', 'region_match': True, 'region_reason': 'Регион'}
    assert parse_analysis(json.dumps(response), TEXT)['relevance'] == 'review'
    response['evidence'] = TEXT[:60]
    assert parse_analysis(json.dumps(response), TEXT)['relevance'] == 'relevant'
    response['region_match'] = None
    assert parse_analysis(json.dumps(response), TEXT)['relevance'] == 'review'
    response.update(region_match=True, evidence='В Красноярске завершено благоустройство набережной Енисея.')
    with_newlines = TEXT.replace('завершено благоустройство', 'завершено\nблагоустройство')
    verified = parse_analysis(json.dumps(response), with_newlines)
    assert verified['relevance'] == 'relevant'
    assert verified['evidence'] in with_newlines
    with pytest.raises(ServiceError):
        parse_analysis('not json', TEXT)
    response['evidence_id'] = 2
    assert parse_analysis(json.dumps(response), TEXT, {2: TEXT[:70]})['evidence'] == TEXT[:70]
    response['evidence_id'] = 999
    assert parse_analysis(json.dumps(response), TEXT, {2: TEXT[:70]})['relevance'] == 'review'


def test_original_extraction_and_login_rejection():
    html = f'<html><head><title>Новость</title><meta property="article:published_time" content="2026-10-12T10:00:00+07:00"></head><body><article>{TEXT}</article></body></html>'
    result = extract(html, 'https://example.com/article')
    assert result['published_at'] == '2026-10-12T10:00:00+07:00'
    assert result['original_text'] == TEXT
    structured = f'<meta itemprop="datePublished" content="2025-12-19T10:16:00+07:00"><div class="wcol article">{TEXT}</div>'
    assert extract(structured, 'https://example.com/news')['published_at'] == '2025-12-19T10:16:00+07:00'
    local = f'<div class="di3-meta__date">23.04.2025</div><article><div class="js-mediator-article">{TEXT}</div>Footer unrelated</article>'
    extracted = extract(local, 'https://example.com/article')
    assert extracted['published_at'] == '2025-04-23'
    assert 'Footer unrelated' not in extracted['original_text']
    with pytest.raises(ValueError):
        extract('<title>Вход</title><meta property="og:description" content="Текст новости"><form>Войти</form>', 'https://example.com')


@pytest.mark.parametrize('url', ['http://localhost/x', 'http://127.0.0.1/a', 'http://10.0.0.1/a', 'file:///etc/passwd', 'https://user:pass@example.com/', 'https://example.com:8000/'])
def test_unsafe_urls_rejected(url):
    with pytest.raises(ValueError):
        normalize_url(url)


def test_dns_rebinding_and_normalization(monkeypatch):
    monkeypatch.setattr('socket.getaddrinfo', lambda *a, **kw: [(2, 1, 6, '', ('127.0.0.1', 443))])
    with pytest.raises(ValueError):
        public_addresses('public.example.com', 443)
    assert normalize_url('https://m.vk.ru/feed?w=wall-123_456&utm_source=a') == 'https://vk.com/wall-123_456'
    assert source_matches('https://example.com/group/1/posts/2', 'https://example.com/group/1')
    assert not source_matches('https://example.com/group/12/posts', 'https://example.com/group/1')


def test_yandex_xml_response():
    xml = b'<yandexsearch><response><results><grouping><group><doc><url>https://example.com/post</url><title>A <hlword>title</hlword></title><passages><passage>Text</passage></passages></doc></group></grouping></results></response></yandexsearch>'
    result = parse_search_response({'rawData': base64.b64encode(xml).decode()})
    assert result[0]['title'] == 'A title'
    assert result[0]['snippet'] == 'Text'
    with pytest.raises(ServiceError):
        parse_search_response({'rawData': 'bad'})
